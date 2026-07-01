"""Shared export utilities (Excel + CSV) for HMS reports."""
import csv
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_DARK_BLUE  = '1E3A5F'
_LIGHT_GRAY = 'F7F9FC'
_ALT_GRAY   = 'EDF2F7'


def excel_response(filename: str) -> HttpResponse:
    r = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    r['Content-Disposition'] = f'attachment; filename="{filename}"'
    return r


def csv_response(filename: str) -> HttpResponse:
    r = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    r['Content-Disposition'] = f'attachment; filename="{filename}"'
    return r


def build_workbook(title: str, period: str, generated_by: str,
                   headers: list[str], rows, sheet_name: str = 'Report') -> Workbook:
    """Return a styled Workbook with a 4-row hospital header + data table."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    n_cols = max(len(headers), 1)
    last_col = get_column_letter(n_cols)

    meta = [
        ('HOSPITAL MANAGEMENT SYSTEM',     13, True,  _DARK_BLUE),
        (title,                            11, True,  _DARK_BLUE),
        (f'Period: {period}',               9, False, '444444'),
        (f'Generated: {timezone.localtime().strftime("%d %b %Y  %H:%M")}  |  '
         f'By: {generated_by}',            8, False, '666666'),
    ]
    for row_no, (text, size, bold, color) in enumerate(meta, 1):
        ws.merge_cells(f'A{row_no}:{last_col}{row_no}')
        cell = ws.cell(row=row_no, column=1, value=text)
        cell.font = Font(bold=bold, size=size, color=color)
        cell.alignment = Alignment(horizontal='center')
        ws.row_dimensions[row_no].height = 16 + (row_no == 1) * 4

    ws.append([])  # row 5 spacer

    # Header row
    for c_idx, hdr in enumerate(headers, 1):
        cell = ws.cell(row=6, column=c_idx, value=hdr)
        cell.font = Font(bold=True, color='FFFFFF', size=10)
        cell.fill = PatternFill(fgColor=_DARK_BLUE, fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[6].height = 18

    # Data rows
    last_data_row = 6
    for r_idx, row in enumerate(rows, 7):
        last_data_row = r_idx
        fill_color = _ALT_GRAY if r_idx % 2 == 0 else _LIGHT_GRAY
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.alignment = Alignment(vertical='center')
            cell.fill = PatternFill(fgColor=fill_color, fill_type='solid')

    # Auto-width (approximate)
    for c_idx, hdr in enumerate(headers, 1):
        col_letter = get_column_letter(c_idx)
        ws.column_dimensions[col_letter].width = max(len(str(hdr)) + 4, 10)

    ws.freeze_panes = 'A7'
    return wb


def send_workbook(wb: Workbook, response: HttpResponse) -> HttpResponse:
    buf = BytesIO()
    wb.save(buf)
    response.write(buf.getvalue())
    return response


def write_csv(response: HttpResponse, headers: list[str], rows) -> HttpResponse:
    writer = csv.writer(response)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(['' if v is None else v for v in row])
    return response
