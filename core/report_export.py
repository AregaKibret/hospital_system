"""Shared Excel export helper for report views.

Pairs with core/templates/reports/_toolbar.html, which links to
`?...&export=excel` and `?...&export=csv` on every report page. Reports
handle `export=csv` themselves (trivial with the stdlib csv module); this
module gives them a one-line way to also handle `export=excel` without each
view hand-rolling openpyxl formatting.
"""
from django.http import HttpResponse


def export_excel(filename, headers, rows, title=''):
    """Return an HttpResponse containing a formatted .xlsx file.

    `headers` is a list of column titles. `rows` is an iterable of
    same-length iterables (e.g. lists/tuples) of cell values.
    """
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (title or 'Report')[:31]

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1E3A5F', end_color='1E3A5F', fill_type='solid')

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 18

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
